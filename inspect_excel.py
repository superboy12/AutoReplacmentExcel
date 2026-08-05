"""
inspect_excel.py
----------------
Script bantu untuk membaca struktur file Excel:
- Menampilkan nama semua sheet
- Menampilkan isi 20 baris pertama setiap sheet (non-kosong)

Ubah FILE_PATH ke salah satu file Excel yang ingin Anda periksa.
Jalankan: python inspect_excel.py
"""

from openpyxl import load_workbook

# ======================================================
# Ubah path di bawah ke salah satu file Excel Anda
FILE_PATH = r"C:\Excel\Data\contoh.xlsx"
# ======================================================

wb = load_workbook(FILE_PATH, data_only=True)

print(f"\n{'='*65}")
print(f"  File   : {FILE_PATH}")
print(f"  Sheet  : {wb.sheetnames}")
print(f"{'='*65}\n")

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n  ── SHEET: [{sname}]  (max_row={ws.max_row}, max_col={ws.max_column}) ──")
    printed = 0
    for row in ws.iter_rows():
        row_vals = [(cell.coordinate, cell.value) for cell in row if cell.value is not None]
        if row_vals:
            for coord, val in row_vals:
                short_val = str(val)[:80]
                print(f"    {coord:>6} │ {short_val}")
            printed += 1
        if printed >= 20:
            print("    ... (baris selanjutnya tidak ditampilkan)")
            break

print(f"\n{'='*65}\n")
